from __future__ import annotations

import json
import zipfile
from pathlib import Path
from typing import Any

from app.services.chunking import format_timestamp
from app.services.domain import Transcript


def _srt_timestamp(seconds: float) -> str:
    milliseconds = max(0, int(round(seconds * 1000)))
    hours, rem = divmod(milliseconds, 3_600_000)
    minutes, rem = divmod(rem, 60_000)
    secs, millis = divmod(rem, 1000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d},{millis:03d}"


def _vtt_timestamp(seconds: float) -> str:
    return _srt_timestamp(seconds).replace(",", ".")


def export_all(
    output_dir: Path,
    source_url: str,
    platform: str,
    title: str,
    metadata: dict[str, Any],
    transcript: Transcript,
    summary: str | None,
) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    plain_text = "\n".join(segment.text for segment in transcript.segments).strip() + "\n"
    txt_path = output_dir / "transcript.txt"
    txt_path.write_text(plain_text, encoding="utf-8")

    srt_lines: list[str] = []
    for index, segment in enumerate(transcript.segments, start=1):
        srt_lines.extend(
            [
                str(index),
                f"{_srt_timestamp(segment.start)} --> {_srt_timestamp(segment.end)}",
                segment.text,
                "",
            ]
        )
    srt_path = output_dir / "subtitles.srt"
    srt_path.write_text("\n".join(srt_lines), encoding="utf-8")

    vtt_lines = ["WEBVTT", ""]
    for segment in transcript.segments:
        vtt_lines.extend([f"{_vtt_timestamp(segment.start)} --> {_vtt_timestamp(segment.end)}", segment.text, ""])
    vtt_path = output_dir / "subtitles.vtt"
    vtt_path.write_text("\n".join(vtt_lines), encoding="utf-8")

    payload = {
        "source_url": source_url,
        "platform": platform,
        "title": title,
        "metadata": metadata,
        "transcript": {
            "language": transcript.language,
            "language_probability": transcript.language_probability,
            "duration": transcript.duration,
            "segments": [
                {"start": item.start, "end": item.end, "text": item.text}
                for item in transcript.segments
            ],
        },
        "summary": summary,
    }
    json_path = output_dir / "result.json"
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    transcript_md = "\n".join(
        f"- **{format_timestamp(item.start)}** {item.text}" for item in transcript.segments
    )
    summary_md = summary or "_未生成 AI 摘要（已关闭 AI 摘要或尚未配置 API Key）。_"
    source_display = "本地上传文件" if platform == "local" else source_url
    markdown = (
        f"# {title}\n\n"
        f"- 平台：`{platform}`\n"
        f"- 来源：{source_display}\n"
        f"- 语言：`{transcript.language}`\n"
        f"- 时长：`{format_timestamp(transcript.duration)}`\n\n"
        "## AI 摘要\n\n"
        f"{summary_md}\n\n"
        "## 完整文字稿\n\n"
        f"{transcript_md}\n"
    )
    md_path = output_dir / "result.md"
    md_path.write_text(markdown, encoding="utf-8")

    # DOCX is intentionally generated from the same canonical transcript and
    # summary, so all downloadable formats stay consistent after an edit.
    from docx import Document

    document = Document()
    document.add_heading(title, level=0)
    document.add_paragraph(f"平台：{platform}    时长：{format_timestamp(transcript.duration)}")
    document.add_heading("AI 摘要", level=1)
    document.add_paragraph(summary or "未生成 AI 摘要。")
    document.add_heading("完整文字稿", level=1)
    for item in transcript.segments:
        document.add_paragraph(f"[{format_timestamp(item.start)}] {item.text}")
    docx_path = output_dir / "result.docx"
    document.save(docx_path)

    zip_path = output_dir / "all-files.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in (md_path, txt_path, srt_path, vtt_path, json_path, docx_path):
            archive.write(path, path.name)

    return {"markdown": md_path, "txt": txt_path, "srt": srt_path, "vtt": vtt_path, "json": json_path, "docx": docx_path, "zip": zip_path}
